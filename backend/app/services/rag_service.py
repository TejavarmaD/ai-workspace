import os
import uuid
import re
from typing import Optional
from sqlalchemy.orm import Session
from fastapi import HTTPException
from backend.app.db.models.knowledge import KnowledgeBase, KnowledgeDocument
from backend.app.db.models.workspace import WorkspaceMember
from backend.app.core.logging import get_logger

logger = get_logger(__name__)

CHROMA_PATH = "/tmp/axiom_chroma"


def get_chroma_client():
    import chromadb
    return chromadb.PersistentClient(path=CHROMA_PATH)


def get_embedding_function():
    from chromadb.utils import embedding_functions
    return embedding_functions.DefaultEmbeddingFunction()


def verify_workspace_access(db: Session, workspace_id: uuid.UUID, user_id: uuid.UUID):
    member = db.query(WorkspaceMember).filter(
        WorkspaceMember.workspace_id == workspace_id,
        WorkspaceMember.user_id == user_id,
        WorkspaceMember.status == "active",
    ).first()
    if not member:
        raise HTTPException(status_code=403, detail="Not a member of this workspace")
    return member


def create_knowledge_base(
    db: Session,
    user_id: uuid.UUID,
    workspace_id: uuid.UUID,
    name: str,
    description: str = None,
) -> KnowledgeBase:
    verify_workspace_access(db, workspace_id, user_id)

    # Create unique collection name
    safe_name = re.sub(r'[^a-z0-9-]', '-', name.lower().strip())
    collection_name = f"kb-{str(uuid.uuid4())[:8]}-{safe_name[:20]}"

    kb = KnowledgeBase(
        workspace_id=workspace_id,
        user_id=user_id,
        name=name,
        description=description,
        collection_name=collection_name,
        document_count=0,
    )
    db.add(kb)
    db.commit()
    db.refresh(kb)

    # Create ChromaDB collection
    try:
        client = get_chroma_client()
        client.create_collection(
            name=collection_name,
            embedding_function=get_embedding_function(),
        )
        logger.info("ChromaDB collection created", collection=collection_name)
    except Exception as e:
        logger.error("Failed to create ChromaDB collection", error=str(e))

    return kb


def get_knowledge_bases(
    db: Session,
    user_id: uuid.UUID,
    workspace_id: uuid.UUID,
) -> list:
    verify_workspace_access(db, workspace_id, user_id)
    return db.query(KnowledgeBase).filter(
        KnowledgeBase.workspace_id == workspace_id,
        KnowledgeBase.user_id == user_id,
    ).order_by(KnowledgeBase.created_at.desc()).all()


def delete_knowledge_base(
    db: Session,
    kb_id: uuid.UUID,
    user_id: uuid.UUID,
) -> None:
    kb = db.query(KnowledgeBase).filter(KnowledgeBase.id == kb_id).first()
    if not kb:
        raise HTTPException(status_code=404, detail="Knowledge base not found")
    if kb.user_id != user_id:
        raise HTTPException(status_code=403, detail="Access denied")

    # Delete ChromaDB collection
    try:
        client = get_chroma_client()
        client.delete_collection(kb.collection_name)
    except Exception as e:
        logger.warning("Failed to delete ChromaDB collection", error=str(e))

    db.delete(kb)
    db.commit()


def extract_text_from_file(file_content: bytes, filename: str) -> str:
    ext = filename.lower().split('.')[-1]

    if ext == 'txt' or ext == 'md':
        return file_content.decode('utf-8', errors='ignore')

    elif ext == 'pdf':
        try:
            from pypdf import PdfReader
            import io
            reader = PdfReader(io.BytesIO(file_content))
            text = ""
            for page in reader.pages:
                text += page.extract_text() + "\n"
            return text
        except Exception as e:
            logger.error("PDF extraction failed", error=str(e))
            return ""

    elif ext in ['docx', 'doc']:
        try:
            from docx import Document
            import io
            doc = Document(io.BytesIO(file_content))
            return "\n".join([p.text for p in doc.paragraphs])
        except Exception as e:
            logger.error("DOCX extraction failed", error=str(e))
            return ""

    else:
        try:
            return file_content.decode('utf-8', errors='ignore')
        except:
            return ""


def chunk_text(text: str, chunk_size: int = 500, overlap: int = 50) -> list[str]:
    words = text.split()
    chunks = []
    i = 0
    while i < len(words):
        chunk = ' '.join(words[i:i + chunk_size])
        if chunk.strip():
            chunks.append(chunk)
        i += chunk_size - overlap
    return chunks


def add_document_to_kb(
    db: Session,
    kb_id: uuid.UUID,
    user_id: uuid.UUID,
    filename: str,
    file_content: bytes,
    file_type: str,
) -> KnowledgeDocument:
    kb = db.query(KnowledgeBase).filter(KnowledgeBase.id == kb_id).first()
    if not kb:
        raise HTTPException(status_code=404, detail="Knowledge base not found")
    if kb.user_id != user_id:
        raise HTTPException(status_code=403, detail="Access denied")

    doc = KnowledgeDocument(
        knowledge_base_id=kb_id,
        filename=filename,
        file_type=file_type,
        file_size=len(file_content),
        status="processing",
        chunk_count=0,
    )
    db.add(doc)
    db.commit()
    db.refresh(doc)

    try:
        text = extract_text_from_file(file_content, filename)
        if not text.strip():
            raise ValueError("No text could be extracted from this file")

        chunks = chunk_text(text)
        if not chunks:
            raise ValueError("No chunks created from extracted text")

        client = get_chroma_client()
        collection = client.get_collection(
            name=kb.collection_name,
            embedding_function=get_embedding_function(),
        )

        chunk_ids = [f"{str(doc.id)}-chunk-{i}" for i in range(len(chunks))]
        metadatas = [{"doc_id": str(doc.id), "filename": filename, "chunk_index": i}
                     for i in range(len(chunks))]

        collection.add(
            documents=chunks,
            ids=chunk_ids,
            metadatas=metadatas,
        )

        doc.status = "ready"
        doc.chunk_count = len(chunks)
        kb.document_count = db.query(KnowledgeDocument).filter(
            KnowledgeDocument.knowledge_base_id == kb_id,
            KnowledgeDocument.status == "ready",
        ).count() + 1

        db.commit()
        logger.info("Document added to KB", doc_id=str(doc.id), chunks=len(chunks))

    except Exception as e:
        doc.status = "failed"
        doc.error_message = str(e)[:500]
        db.commit()
        logger.error("Document processing failed", error=str(e))

    db.refresh(doc)
    return doc


def search_knowledge_base(
    kb_collection_name: str,
    query: str,
    n_results: int = 5,
) -> list[str]:
    try:
        client = get_chroma_client()
        collection = client.get_collection(
            name=kb_collection_name,
            embedding_function=get_embedding_function(),
        )
        results = collection.query(
            query_texts=[query],
            n_results=min(n_results, collection.count()),
        )
        if results and results.get('documents'):
            return results['documents'][0]
        return []
    except Exception as e:
        logger.error("Knowledge base search failed", error=str(e))
        return []


def build_rag_context(
    db: Session,
    kb_id: uuid.UUID,
    user_id: uuid.UUID,
    query: str,
    n_results: int = 5,
) -> Optional[str]:
    kb = db.query(KnowledgeBase).filter(KnowledgeBase.id == kb_id).first()
    if not kb or kb.user_id != user_id:
        return None

    chunks = search_knowledge_base(kb.collection_name, query, n_results)
    if not chunks:
        return None

    context = f"Relevant information from knowledge base '{kb.name}':\n\n"
    for i, chunk in enumerate(chunks, 1):
        context += f"[{i}] {chunk}\n\n"

    return context


def format_kb(kb: KnowledgeBase) -> dict:
    return {
        "id": str(kb.id),
        "name": kb.name,
        "description": kb.description,
        "collection_name": kb.collection_name,
        "document_count": kb.document_count,
        "workspace_id": str(kb.workspace_id),
        "user_id": str(kb.user_id),
        "created_at": str(kb.created_at),
    }


def format_document(doc: KnowledgeDocument) -> dict:
    return {
        "id": str(doc.id),
        "filename": doc.filename,
        "file_type": doc.file_type,
        "file_size": doc.file_size,
        "chunk_count": doc.chunk_count,
        "status": doc.status,
        "error_message": doc.error_message,
        "created_at": str(doc.created_at),
    }